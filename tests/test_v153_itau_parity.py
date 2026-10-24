"""
test_v153_itau_parity.py
========================
Suíte de Testes Automatizados para a Release v1.53:
"Paridade Estrutural e de Categorização do Extrato Itaú"
Cobrindo Histórias de Usuário BDD: US-01, US-02, US-03, US-04 e US-05.
"""

import io
import os
import sys
import unittest
from unittest.mock import MagicMock, patch
import pandas as pd
import openpyxl

# Assegura que o diretório raiz do backend esteja no sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from src.utils.itau_parser import processar_extrato_itau_bytes, limpar_valor
from src.services.extrato_service import _salvar_transacoes, _parse_date


class TestV153ItauParity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Amostra real do extrato Itaú (Setembro/2026) em CSV com delimitador vírgula e decimal ponto
        cls.csv_amostra_real = (
            "data,lançamento,ag./origem,valor (R$),saldo (R$)\n"
            "31/08/2026,SALDO ANTERIOR,,,46085.89\n"
            "01/09/2026,BOLETO PAGO ENEL,, -4216.93, 41868.96\n"
            "01/09/2026,SALDO TOTAL DISPONÍVEL DIA,,, 41868.96\n"
            "01/09/2026,BOLETOS RECEBIDOS 01/09S,, 2500.00, 44368.96\n"
            "02/09/2026,REND PAGO APLIC AUT MAIS,, 124.50, 44493.46\n"
            "02/09/2026,APLICACAO CDB PLUS,, -5000.00, 39493.46\n"
            "03/09/2026,RESG AUT APLIC INVEST,, 3000.00, 42493.46\n"
            "03/09/2026,TAR/CUSTAS COBRANCA,, -35.80, 42457.66\n"
            "04/09/2026,PIX ENVIADO SILVANDIRA D,, -150.00, 42307.66\n"
            "04/09/2026,SISPAG CONCESSIONARIA,, -820.40, 41487.26\n"
            "05/09/2026,RSHOP CIDEAALMADA-001032,, -45.90, 41441.36\n"
            "05/09/2026,SALDO TOTAL DISPONÍVEL DIA,,, 41441.36\n"
        ).encode('utf-8')

    # =========================================================================
    # US-01: Ingestão Resiliente Multi-formato (XLSX, XLS, CSV Multi-encoding)
    # =========================================================================
    def test_us01_cenario01_ingestao_csv_virgula_e_decimal_ponto(self):
        """Cenário 01: Ingestão de CSV Itaú com delimitador vírgula e decimal com ponto."""
        df, excel_io, nome_saida = processar_extrato_itau_bytes(
            self.csv_amostra_real,
            nome_arquivo="extrato_itau_setembro.csv"
        )
        self.assertFalse(df.empty)
        self.assertEqual(nome_saida, "extrato_itau_setembro_PROC.xlsx")

        # Verifica se o boleto Enel foi extraído com débito 4216.93 e data 01/09/2026
        linha_enel = df[df['Lançamento'].str.contains("BOLETO PAGO ENEL", case=False, na=False)]
        self.assertEqual(len(linha_enel), 1)
        self.assertEqual(linha_enel.iloc[0]['Débito (R$)'], 4216.93)
        self.assertEqual(linha_enel.iloc[0]['Data_Valida'], "01/09/2026")

    def test_us01_cenario02_ingestao_csv_ponto_e_virgula_e_latin1(self):
        """Cenário 02: Ingestão de CSV Itaú com delimitador ponto-e-vírgula e encoding Latin-1."""
        csv_latin1 = (
            "Data;Lançamento;Crédito (R$);Débito (R$)\n"
            "31/08/2026;SALDO ANTERIOR;10000,00;\n"
            "01/09/2026;LIQUIDAÇÃO DE COBRANÇA;1.250,50;\n"
            "02/09/2026;BOLETO PAGO FORNECEDOR;;540,20\n"
        ).encode('latin-1')

        df, excel_io, nome_saida = processar_extrato_itau_bytes(csv_latin1, "itau_latin1.csv")
        self.assertFalse(df.empty)
        self.assertEqual(nome_saida, "itau_latin1_PROC.xlsx")

        # Valida preservação de acentuação e conversão numérica
        linha_liq = df[df['Lançamento'].str.contains("LIQUIDAÇÃO DE COBRANÇA", case=False, na=False)]
        self.assertEqual(len(linha_liq), 1)
        self.assertEqual(linha_liq.iloc[0]['Crédito (R$)'], 1250.50)

    def test_us01_cenario03_ingestao_nativa_xlsx(self):
        """Cenário 03: Ingestão nativa de pasta de trabalho XLSX."""
        dados = {
            "Data": ["01/09/2026", "02/09/2026"],
            "Lançamento": ["BOLETOS RECEBIDOS", "TAR/CUSTAS COBRANCA"],
            "Valor": [1500.00, -25.50],
        }
        df_in = pd.DataFrame(dados)
        buf = io.BytesIO()
        df_in.to_excel(buf, index=False)
        buf.seek(0)

        df, excel_io, nome_saida = processar_extrato_itau_bytes(buf.getvalue(), "extrato_itau.xlsx")
        self.assertEqual(len(df), 2)
        self.assertEqual(nome_saida, "extrato_itau_PROC.xlsx")
        self.assertEqual(df.iloc[0]['Crédito (R$)'], 1500.00)
        self.assertEqual(df.iloc[1]['Débito (R$)'], 25.50)

    def test_us01_cenario04_excecao_graciosa_arquivo_invalido(self):
        """Cenário 04: Tratamento de exceção graciosa para arquivos corrompidos ou ilegíveis."""
        conteudo_corrompido = b"\x00\x01\x02\x03\x04\x05"
        with self.assertRaises(ValueError) as ctx:
            processar_extrato_itau_bytes(conteudo_corrompido, "corrompido.xlsx")
        self.assertIn("Não foi possível ler o arquivo de extrato do Itaú como Excel ou CSV.", str(ctx.exception))

    # =========================================================================
    # US-02: Extração de Saldo Anterior e Expurgo de Linhas Transitórias
    # =========================================================================
    def test_us02_cenario01_extracao_saldo_anterior(self):
        """Cenário 01: Extração precisa de Saldo Anterior decimal em extrato Itaú."""
        df, excel_io, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        wb = openpyxl.load_workbook(excel_io)
        ws_cons = wb['Consolidado']

        # O saldo anterior extraído de 31/08/2026,SALDO ANTERIOR,,,46085.89 deve estar na célula E2
        self.assertEqual(ws_cons['A2'].value, "Saldo Anterior")
        self.assertAlmostEqual(float(ws_cons['E2'].value), 46085.89, places=2)

    def test_us02_cenario02_remocao_saldo_anterior_do_dataframe(self):
        """Cenário 02: Remoção mandatória da linha de Saldo Anterior do DataFrame analítico."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        self.assertFalse(any(df['Lançamento'].str.contains("SALDO ANTERIOR", case=False, na=False)))

    def test_us02_cenario03_descarte_saldos_diarios_transitorios(self):
        """Cenário 03: Descarte completo de ocorrências de 'SALDO TOTAL DISPONÍVEL DIA' e termos transitórios."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        self.assertFalse(any(df['Lançamento'].str.contains("SALDO TOTAL DISPONÍVEL", case=False, na=False)))
        self.assertFalse(any(df['Lançamento'].str.contains("SALDO", case=False, na=False)))
        self.assertFalse(any(df['Lançamento'].str.contains("TOTAL", case=False, na=False)))

    def test_us02_cenario04_filtro_competencia_principal(self):
        """Cenário 04: Filtro de competência principal descartando resquícios de outros meses."""
        # Se houvesse transação de 31/08/2026 diferente de saldo anterior, seria expurgada pela moda
        csv_multi_mes = (
            "data,lançamento,valor (R$)\n"
            "31/08/2026,LANCAMENTO RESIDUAL AGOSTO,-100.00\n"
            "01/09/2026,BOLETO PAGO ENEL,-500.00\n"
            "02/09/2026,BOLETO PAGO AGUA,-300.00\n"
            "03/09/2026,BOLETO PAGO GAS,-200.00\n"
        ).encode('utf-8')

        df, _, _ = processar_extrato_itau_bytes(csv_multi_mes, "itau.csv")
        self.assertEqual(len(df), 3)
        self.assertTrue(all(df['Data_Valida'].str.endswith("/09/2026")))

    # =========================================================================
    # US-03: Taxonomia Granular de 7 Categorias Contábeis
    # =========================================================================
    def test_us03_cenario01_classificacao_receitas(self):
        """Cenário 01: Classificação de receitas de arrecadação e boletos recebidos."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        row = df[df['Lançamento'].str.contains("BOLETOS RECEBIDOS", case=False, na=False)].iloc[0]
        self.assertEqual(row['Categoria'], 'Receitas')

    def test_us03_cenario02_classificacao_aplicacoes_e_resgates_separadas(self):
        """Cenário 02: Classificação de aplicações e resgates em categorias separadas e distintas."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        row_aplic = df[df['Lançamento'].str.contains("APLICACAO CDB PLUS", case=False, na=False)].iloc[0]
        row_resg = df[df['Lançamento'].str.contains("RESG AUT APLIC INVEST", case=False, na=False)].iloc[0]

        self.assertEqual(row_aplic['Categoria'], 'Aplicações')
        self.assertEqual(row_resg['Categoria'], 'Resgates')
        self.assertFalse(any(df['Categoria'] == 'Aplicações/Resgates'))

    def test_us03_cenario03_classificacao_rentabilidades(self):
        """Cenário 03: Classificação de rendimentos financeiros como Rentabilidades."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        row = df[df['Lançamento'].str.contains("REND PAGO", case=False, na=False)].iloc[0]
        self.assertEqual(row['Categoria'], 'Rentabilidades')

    def test_us03_cenario04_classificacao_tarifas(self):
        """Cenário 04: Classificação de custas e tarifas bancárias como Tarifas."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        row = df[df['Lançamento'].str.contains("TAR/CUSTAS", case=False, na=False)].iloc[0]
        self.assertEqual(row['Categoria'], 'Tarifas')

    def test_us03_cenario05_classificacao_outros_gastos_operacionais(self):
        """Cenário 05: Classificação de pagamentos operacionais (SISPAG, RSHOP, PIX, Boleto Pago) como Outros Gastos."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")
        termos_gastos = [
            "BOLETO PAGO ENEL",
            "PIX ENVIADO SILVANDIRA",
            "SISPAG CONCESSIONARIA",
            "RSHOP CIDEAALMADA"
        ]
        for termo in termos_gastos:
            subset = df[df['Lançamento'].str.contains(termo, case=False, na=False)]
            self.assertEqual(len(subset), 1, f"Termo {termo} não encontrado.")
            self.assertEqual(subset.iloc[0]['Categoria'], 'Outros Gastos', f"Termo {termo} não categorizado como Outros Gastos.")

    # =========================================================================
    # US-04: Geração de Planilha XLSX com Duas Abas e Fórmulas Dinâmicas
    # =========================================================================
    def test_us04_cenario01_workbook_abas_detalhado_e_consolidado(self):
        """Cenário 01: Criação de Workbook com exatamente as abas 'Detalhado' e 'Consolidado'."""
        _, excel_io, nome_saida = processar_extrato_itau_bytes(self.csv_amostra_real, "extrato.csv")
        wb = openpyxl.load_workbook(excel_io)
        self.assertEqual(wb.sheetnames, ['Detalhado', 'Consolidado'])
        self.assertEqual(nome_saida, "extrato_PROC.xlsx")

    def test_us04_cenario02_formulas_subtotal_na_aba_detalhado(self):
        """Cenário 02: Injeção de fórmulas nativas de subtotal =SUM(...) na aba Detalhado."""
        _, excel_io, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "extrato.csv")
        wb = openpyxl.load_workbook(excel_io, data_only=False)
        ws_det = wb['Detalhado']

        # Encontra as linhas de subtotal "Total Outros Gastos"
        formulas_encontradas = []
        for row in ws_det.iter_rows(values_only=False):
            cell_a = row[0]
            if cell_a.value and str(cell_a.value).startswith("Total "):
                # Coluna C (Crédito) e Coluna D (Débito)
                val_c = str(row[2].value or "")
                val_d = str(row[3].value or "")
                if "=SUM" in val_c or "=SUM" in val_d:
                    formulas_encontradas.append((cell_a.value, val_c, val_d))

        self.assertGreater(len(formulas_encontradas), 0)
        for cat_label, f_cred, f_deb in formulas_encontradas:
            self.assertTrue(f_cred.startswith("=SUM(C") or f_deb.startswith("=SUM(D"))

    def test_us04_cenario03_aba_consolidado_saldo_anterior_e_vinculos(self):
        """Cenário 03: Composição da aba Consolidado partindo de Saldo Anterior e fórmulas vinculadas."""
        _, excel_io, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "extrato.csv")
        wb = openpyxl.load_workbook(excel_io, data_only=False)
        ws_cons = wb['Consolidado']

        # Cabeçalhos
        self.assertEqual(ws_cons['A1'].value, "Categoria")
        self.assertEqual(ws_cons['D1'].value, "Crédito (R$)")
        self.assertEqual(ws_cons['E1'].value, "Débito (R$)")

        # Saldo Anterior na linha 2
        self.assertEqual(ws_cons['A2'].value, "Saldo Anterior")
        self.assertAlmostEqual(float(ws_cons['E2'].value), 46085.89, places=2)

        # Categorias vinculadas
        linhas_vinculadas = []
        for r in range(3, ws_cons.max_row):
            cat_name = ws_cons.cell(row=r, column=1).value
            form_cred = str(ws_cons.cell(row=r, column=4).value or "")
            form_deb = str(ws_cons.cell(row=r, column=5).value or "")
            if cat_name and cat_name != "Saldo Final":
                linhas_vinculadas.append((cat_name, form_cred, form_deb))

        self.assertGreater(len(linhas_vinculadas), 0)
        for cat_name, form_cred, form_deb in linhas_vinculadas:
            self.assertTrue(form_cred.startswith("=Detalhado!C"), f"Fórmula de crédito incorreta para {cat_name}: {form_cred}")
            self.assertTrue(form_deb.startswith("=Detalhado!D"), f"Fórmula de débito incorreta para {cat_name}: {form_deb}")

    def test_us04_cenario04_formula_saldo_final(self):
        """Cenário 04: Fórmula e cálculo dinâmico de Saldo Final fechando conciliação contábil."""
        _, excel_io, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "extrato.csv")
        wb = openpyxl.load_workbook(excel_io, data_only=False)
        ws_cons = wb['Consolidado']

        # Última linha é Saldo Final
        max_row = ws_cons.max_row
        self.assertEqual(ws_cons.cell(row=max_row, column=1).value, "Saldo Final")

        formula_final = str(ws_cons.cell(row=max_row, column=5).value)
        n = max_row - 1
        expected_formula = f"=E2 + (SUM(D3:D{n}) - (SUM(E3:E{n})*-1))"
        self.assertEqual(formula_final, expected_formula)

    # =========================================================================
    # US-05: Blindagem da Conciliação Bancária (_salvar_transacoes)
    # =========================================================================
    def test_us05_cenario01_insercao_exclusiva_outros_gastos(self):
        """Cenário 01: Inserção exclusiva de transações 'Outros Gastos' no Supabase."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")

        # Mock do cliente Supabase para testar a persistência
        mock_supabase = MagicMock()
        mock_table = MagicMock()
        mock_supabase.table.return_value = mock_table
        mock_table.insert.return_value = mock_table

        with patch("src.services.extrato_service._get_supabase", return_value=mock_supabase):
            qtd = _salvar_transacoes(
                df=df,
                banco="itau",
                storage_path="condo-01/extrato.csv",
                administradora_id="adm-alpha",
                condominio_id="condo-alpha-01",
                fonte_pagadora_id="fp-123"
            )

        # Na amostra real temos exatamente 4 registros em Outros Gastos:
        # BOLETO PAGO ENEL, PIX ENVIADO SILVANDIRA, SISPAG CONCESSIONARIA, RSHOP CIDEAALMADA
        self.assertEqual(qtd, 4)
        args, kwargs = mock_table.insert.call_args
        registros_inseridos = args[0]
        self.assertEqual(len(registros_inseridos), 4)

        for reg in registros_inseridos:
            self.assertEqual(reg["metadados"], {"categoria": "Outros Gastos"})
            self.assertEqual(reg["tipo"], "debito")
            self.assertEqual(reg["banco"], "itau")
            self.assertEqual(reg["administradora_id"], "adm-alpha")
            self.assertEqual(reg["condominio_id"], "condo-alpha-01")
            self.assertEqual(reg["fonte_pagadora_id"], "fp-123")

    def test_us05_cenario02_bloqueio_receitas_tarifas_investimentos(self):
        """Cenário 02: Bloqueio estrito de inserção para receitas, investimentos e tarifas."""
        df, _, _ = processar_extrato_itau_bytes(self.csv_amostra_real, "itau.csv")

        # Verifica que o DataFrame possui diversas categorias além de Outros Gastos
        categorias_presentes = set(df['Categoria'].unique())
        self.assertIn("Receitas", categorias_presentes)
        self.assertIn("Rentabilidades", categorias_presentes)
        self.assertIn("Aplicações", categorias_presentes)
        self.assertIn("Resgates", categorias_presentes)
        self.assertIn("Tarifas", categorias_presentes)
        self.assertIn("Outros Gastos", categorias_presentes)

        mock_supabase = MagicMock()
        mock_table = MagicMock()
        mock_supabase.table.return_value = mock_table
        mock_table.insert.return_value = mock_table

        with patch("src.services.extrato_service._get_supabase", return_value=mock_supabase):
            _salvar_transacoes(df, "itau", "caminho")

        args, _ = mock_table.insert.call_args
        inseridos = args[0]
        descricoes = [r["descricao"] for r in inseridos]

        self.assertFalse(any("BOLETOS RECEBIDOS" in d for d in descricoes))
        self.assertFalse(any("REND PAGO" in d for d in descricoes))
        self.assertFalse(any("APLICACAO CDB" in d for d in descricoes))
        self.assertFalse(any("RESG AUT" in d for d in descricoes))
        self.assertFalse(any("TAR/CUSTAS" in d for d in descricoes))

    def test_us05_cenario03_formatacao_data_iso_postgres(self):
        """Cenário 03: Preservação de integridade de data no formato ISO YYYY-MM-DD para o PostgreSQL."""
        data_convertida = _parse_date("15/09/2026")
        self.assertEqual(data_convertida, "2026-09-15")

        data_convertida_iso = _parse_date("2026-09-15")
        self.assertEqual(data_convertida_iso, "2026-09-15")


if __name__ == "__main__":
    unittest.main()
